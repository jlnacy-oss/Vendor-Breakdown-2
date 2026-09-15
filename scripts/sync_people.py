# -*- coding: utf-8 -*-
"""
sync_people.py — add newly researched executives to gip-tracker.xlsx WITHOUT
touching anything you've already typed.

Run this after a research wave lands. It appends a row for any researched person
who isn't in the People sheet yet, and leaves every existing row exactly as it is.
It never edits, reorders or deletes your data.

    python3 scripts/sync_people.py
"""
import os
import sys

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)
sys.path.insert(0, SCRIPT_DIR)

from research_data import VENDORS, person_id

BOOK = os.path.join(ROOT, "gip-tracker.xlsx")

REF_FILL = PatternFill("solid", fgColor="F5F6F8")
REF_FONT = Font(name="Arial", size=10, color="5B6270")
BODY_FONT = Font(name="Arial", size=10, color="14161C")
THIN = Side(style="thin", color="DCDFE4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

ROLE_HINTS = [
    ("chief technology", "CTO"), ("chief executive", "CEO"), ("ceo", "CEO"),
    ("president", "President"), ("chief operating", "COO"),
    ("chief financial", "CFO"), ("chief strategy", "Other"),
]


def role_of(title):
    t = title.lower()
    for hint, role in ROLE_HINTS:
        if hint in t:
            return role
    return "Other"


def main():
    if not os.path.exists(BOOK):
        print("gip-tracker.xlsx not found next to index.html — nothing to sync.")
        return

    wb = load_workbook(BOOK)
    if "People" not in wb.sheetnames:
        print("That workbook has no People sheet — is it the right file?")
        return
    ws = wb["People"]

    existing = set()
    last_row = 1
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(row=r, column=1).value
        if pid:
            existing.add(str(pid).strip())
            last_row = r
        elif any(ws.cell(row=r, column=c).value for c in range(2, 17)):
            last_row = r

    added = 0
    row = last_row + 1
    for v in VENDORS:
        for p in v["people"]:
            pid = person_id(v["slug"], p["name"])
            if pid in existing:
                continue
            title = p["title"].replace("&amp;", "&")
            for ci, val in enumerate([pid, v["name"], p["name"], title, role_of(title)], start=1):
                c = ws.cell(row=row, column=ci, value=val)
                c.font, c.fill, c.border = REF_FONT, REF_FILL, BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top")
            for ci in range(6, 17):
                c = ws.cell(row=row, column=ci)
                c.font, c.border = BODY_FONT, BORDER
                c.alignment = Alignment(wrap_text=True, vertical="top")
            ws.cell(row=row, column=16, value="Yes").font = BODY_FONT
            row += 1
            added += 1

    if not added:
        print("People sheet is already up to date — nothing added.")
        return

    wb.save(BOOK)
    print(f"added {added} new contact row(s) to the People sheet; existing rows untouched")


if __name__ == "__main__":
    main()
