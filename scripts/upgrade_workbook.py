# -*- coding: utf-8 -*-
"""
upgrade_workbook.py — add any missing columns to gip-tracker.xlsx WITHOUT touching
a single value you have already entered.

    python3 scripts/upgrade_workbook.py

Safe to run repeatedly. It only ever appends columns that are absent, adds their
dropdowns, and leaves every existing cell, row and sheet exactly as it was. Run it
after pulling a new version of the app.

Everything it adds to the People sheet is for information YOU collect — typically in
the room, in a meeting. None of it is ever researched or filled automatically; the
privacy guard in people_schema.py blocks personal details from the research layer
entirely.
"""
import os
import shutil
import sys
from datetime import datetime

from openpyxl import load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(SCRIPT_DIR)
BOOK = os.path.join(ROOT, "gip-tracker.xlsx")

HDR_FILL = PatternFill("solid", fgColor="A5702B")
HDR_FONT = Font(name="Arial", size=10, bold=True, color="FFFFFF")
BODY_FONT = Font(name="Arial", size=10, color="14161C")
THIN = Side(style="thin", color="DCDFE4")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# Columns added to the People sheet, in order. (header, width, dropdown-list-or-None)
PEOPLE_COLUMNS = [
    ("Mobile", 18, None),
    ("Preferred Contact", 20, ["Email", "Mobile", "Work phone", "LinkedIn", "Via assistant", "Via colleague"]),
    ("Assistant / EA", 24, None),
    ("Where We Met", 30, None),
    ("Facts From Meetings", 52, None),
    ("Next Touch", 32, None),
    ("Next Touch Due", 16, None),
]

ENGAGEMENT_COLUMNS = [
    ("Contract End", 16, None),
    ("Annual Value Band", 22, ["<$100k", "$100k-$500k", "$500k-$1m", "$1m-$5m", "$5m+", "Unknown"]),
    ("Notice Period", 18, None),
]

DATE_COLUMNS = {"Next Touch Due", "Contract End"}


def headers(ws):
    return [c.value for c in ws[1]]


def add_columns(ws, spec, last_data_row, lists_sheet, list_col_start):
    """Append any missing columns. Returns (added, next_free_list_column)."""
    existing = set(headers(ws))
    added = []
    col = ws.max_column
    list_col = list_col_start

    for header, width, options in spec:
        if header in existing:
            continue
        col += 1
        c = ws.cell(row=1, column=col, value=header)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BORDER
        c.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = width

        for r in range(2, last_data_row + 1):
            cell = ws.cell(row=r, column=col)
            cell.font, cell.border = BODY_FONT, BORDER
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if header in DATE_COLUMNS:
                cell.number_format = "yyyy-mm-dd"

        if options:
            letter = get_column_letter(list_col)
            h = lists_sheet.cell(row=1, column=list_col, value=header)
            h.fill, h.font, h.border = HDR_FILL, HDR_FONT, BORDER
            lists_sheet.column_dimensions[letter].width = max(18, width - 4)
            for i, opt in enumerate(options, start=2):
                oc = lists_sheet.cell(row=i, column=list_col, value=opt)
                oc.font, oc.border = BODY_FONT, BORDER
            dv = DataValidation(type="list",
                                formula1=f"=Lists!${letter}$2:${letter}${1 + len(options)}",
                                allow_blank=True)
            ws.add_data_validation(dv)
            dv.add(f"{get_column_letter(col)}2:{get_column_letter(col)}{last_data_row}")
            list_col += 1

        added.append(header)

    return added, list_col


def main():
    if not os.path.exists(BOOK):
        print("gip-tracker.xlsx not found next to index.html.")
        return 1

    backup = os.path.join(ROOT, f"gip-tracker.backup-{datetime.now():%Y%m%d-%H%M%S}.xlsx")
    shutil.copy2(BOOK, backup)

    wb = load_workbook(BOOK)
    if "Lists" not in wb.sheetnames:
        print("That workbook has no Lists sheet — is it the right file?")
        return 1
    lists = wb["Lists"]
    next_list_col = lists.max_column + 1

    total = []
    for sheet, spec, rows in (("People", PEOPLE_COLUMNS, 399),
                              ("Engagements", ENGAGEMENT_COLUMNS, 399)):
        if sheet not in wb.sheetnames:
            continue
        added, next_list_col = add_columns(wb[sheet], spec, rows, lists, next_list_col)
        if added:
            print(f"{sheet}: added {len(added)} column(s) -> {', '.join(added)}")
            total += added

    if not total:
        os.remove(backup)
        print("Workbook already has every column — nothing changed, no backup needed.")
        return 0

    wb.save(BOOK)
    print(f"\nSaved. Your existing data is untouched.")
    print(f"Backup of the previous version: {os.path.basename(backup)}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
