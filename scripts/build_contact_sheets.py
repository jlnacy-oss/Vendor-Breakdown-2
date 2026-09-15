# -*- coding: utf-8 -*-
"""
build_contact_sheets.py — a one-page meeting-prep sheet for every tracked contact.

    python3 scripts/build_contact_sheets.py              # public sheets (safe to publish)
    python3 scripts/build_contact_sheets.py --with-notes # adds your workbook data,
                                                         # written to private-contact-sheets/
    python3 scripts/build_contact_sheets.py palantir-alex-karp   # one person

The public sheet contains only research-layer material: who the person is, what
their company does, the numbers, where it sits competitively, and what to
pressure-test. It is the sheet you read in the car on the way to the meeting.

The --with-notes build additionally pulls your email, phone, relationship owner,
how you know them, and that vendor's live engagement threads out of
gip-tracker.xlsx. Those files land in private-contact-sheets/, which is listed in
.gitignore and never published.
"""
import os
import sys

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
ROOT_DIR = os.path.dirname(SCRIPT_DIR)
sys.path.insert(0, SCRIPT_DIR)

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from build_pdfs import LINE, INK, MUTED, SAND, STEEL, WHITE, PANEL_BG, _clip, section_head, stat_box, styles
from research_data import person_id
import people_research
from vendor_data import VENDORS as _CORE
from deep_research import DEEP

# The full records (offering, position_note, sentiment, people) live in
# vendor_data.py and the deep_wave modules; research_data.py only carries the
# index view used by the site.
for _v in _CORE:
    _v.setdefault("group", "Physical / Industrial AI")
VENDORS = sorted(_CORE + DEEP, key=lambda v: v["name"].lower())

BUILD_DATE = "September 15, 2026"

styles["mono"] = ParagraphStyle("mono", fontName="Helvetica-Bold", fontSize=20, leading=22,
                                textColor=WHITE, alignment=TA_CENTER)
styles["pname"] = ParagraphStyle("pname", fontName="Helvetica-Bold", fontSize=20, leading=22, textColor=INK)
styles["prole"] = ParagraphStyle("prole", fontName="Helvetica", fontSize=10, leading=12.5, textColor=MUTED)
styles["private_h"] = ParagraphStyle("private_h", fontName="Helvetica-Bold", fontSize=8,
                                     leading=10, textColor=STEEL)


def para(text, style="body"):
    return Paragraph(text, styles[style])


def initials(name):
    parts = [p for p in name.split() if p]
    if not parts:
        return "?"
    return (parts[0][0] + (parts[-1][0] if len(parts) > 1 else "")).upper()


def monogram(name):
    t = Table([[Paragraph(initials(name), styles["mono"])]], colWidths=[0.62 * inch], rowHeights=[0.62 * inch])
    t.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), SAND),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("TOPPADDING", (0, 0), (-1, -1), 0),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 0),
    ]))
    return t


def bullets(items):
    return "<br/>".join("&bull;&nbsp; " + i for i in items)


def build_sheet(person, vendor, out_path, notes=None, engagements=None, level=0, research=None):
    scale = (1.0, 0.75, 0.55, 0.40)[min(level, 3)]

    def clip(text, base):
        return _clip(text, max(80, int(base * scale)))

    doc = SimpleDocTemplate(out_path, pagesize=letter,
                            leftMargin=0.42 * inch, rightMargin=0.42 * inch,
                            topMargin=0.26 * inch, bottomMargin=0.22 * inch)
    story = []
    title = person["title"].replace("&amp;", "&")

    # ---------- header ----------
    head_text = Table(
        [[Paragraph(f"CONTACT ONE-SHEET &nbsp;\u00b7&nbsp; {vendor['name'].upper()} &nbsp;\u00b7&nbsp; {BUILD_DATE.upper()}", styles["kicker"])],
         [Paragraph(person["name"], styles["pname"])],
         [Paragraph(title, styles["prole"])]],
        colWidths=[6.7 * inch])
    head_text.setStyle(TableStyle([
        ("LEFTPADDING", (0, 0), (-1, -1), 8), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
        ("TOPPADDING", (0, 0), (-1, 0), 4), ("BOTTOMPADDING", (0, 0), (-1, 0), 2),
        ("TOPPADDING", (0, 1), (-1, -1), 0), ("BOTTOMPADDING", (0, 1), (-1, -1), 2),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
    ]))
    header = Table([[monogram(person["name"]), head_text]], colWidths=[0.72 * inch, 6.94 * inch])
    header.setStyle(TableStyle([
        ("BACKGROUND", (0, 0), (-1, -1), WHITE),
        ("BOX", (0, 0), (-1, -1), 0.6, LINE),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (0, 0), 5), ("RIGHTPADDING", (0, 0), (0, 0), 0),
        ("TOPPADDING", (0, 0), (-1, -1), 5), ("BOTTOMPADDING", (0, 0), (-1, -1), 5),
    ]))
    story.append(header)
    story.append(Table([[""]], colWidths=[8.5 * inch], rowHeights=[2.4],
                       style=TableStyle([("BACKGROUND", (0, 0), (-1, -1), SAND)])))
    story.append(Spacer(1, 6))

    # ---------- the company, in numbers ----------
    kpis = vendor.get("kpis") or [
        (vendor.get("revenue_short", "\u2014"), "Revenue"),
        (vendor.get("hq", "\u2014").split(",")[0], "HQ"),
    ]
    stats = [stat_box(n, l) for n, l in kpis[:5]]
    n = len(stats)
    row = Table([stats], colWidths=[(7.656 / n) * inch] * n, hAlign="LEFT")
    row.setStyle(TableStyle([("LEFTPADDING", (0, 0), (-1, -1), 2), ("RIGHTPADDING", (0, 0), (-1, -1), 2),
                             ("TOPPADDING", (0, 0), (-1, -1), 0), ("BOTTOMPADDING", (0, 0), (-1, -1), 0)]))
    story.append(row)
    story.append(Spacer(1, 8))

    # ---------- 01 who you're meeting ----------
    story.append(section_head("Who You're Meeting", "01"))
    story.append(Spacer(1, 3))
    story.append(para(f"<b>{person['name']}</b> \u2014 {title}, {vendor['name']}.", "body"))
    if research:
        if research.get("in_role_since"):
            story.append(Spacer(1, 2))
            story.append(para(f"<b>In role:</b> {research['in_role_since']}", "body_sm"))
        story.append(Spacer(1, 2))
        story.append(para(f"<b>Today:</b> {clip(research['focus'], 320)}", "body_sm"))
        story.append(Spacer(1, 2))
        story.append(para(clip(research["background"], 620), "body_sm"))
        if research.get("career"):
            story.append(Spacer(1, 2))
            path = "; ".join(f"{c['role']}, {c['org']}" for c in research["career"][:4])
            story.append(para(f"<b>Career:</b> {clip(path, 360)}", "body_sm"))
        if research.get("education"):
            story.append(Spacer(1, 2))
            story.append(para(f"<b>Education:</b> {clip('; '.join(research['education']), 200)}", "body_sm"))
    elif notes and notes.get("background"):
        story.append(Spacer(1, 2))
        story.append(para(clip(notes["background"], 420), "body_sm"))
    story.append(Spacer(1, 2))
    story.append(para(f"<b>The company:</b> {clip(vendor['description'], 430)}", "body_sm"))
    story.append(Spacer(1, 2))
    story.append(para(f"<b>What they sell:</b> {clip(vendor['offering'], 380)}", "body_sm"))
    story.append(Spacer(1, 7))

    # ---------- 02 where the company sits ----------
    story.append(section_head("Where The Company Sits", "02"))
    story.append(Spacer(1, 3))
    story.append(para(clip(vendor["position_note"], 560), "body_sm"))
    story.append(Spacer(1, 7))

    # ---------- 03 what to pressure-test ----------
    n_bear = 4 if level < 2 else 3
    if research and research.get("talking_points"):
        story.append(section_head("Talking Points", "03"))
        story.append(Spacer(1, 3))
        story.append(para(bullets([clip(t, 200) for t in research["talking_points"][:3]]), "body_sm"))
        story.append(Spacer(1, 5))
        story.append(para("<b>What to pressure-test about the company:</b>", "body_sm"))
        story.append(para(bullets([clip(b, 170) for b in vendor["sentiment_bear"][:max(2, n_bear - 1)]]), "body_sm"))
    else:
        story.append(section_head("What To Pressure-Test", "03"))
        story.append(Spacer(1, 3))
        story.append(para(bullets([clip(b, 190) for b in vendor["sentiment_bear"][:n_bear]]), "body_sm"))
    story.append(Spacer(1, 7))

    # ---------- 04 GIP context ----------
    story.append(section_head("GIP Context", "04"))
    story.append(Spacer(1, 3))
    story.append(para(f"<b>GIP:</b> {clip(vendor['gip_connection'], 340)}", "body_sm"))
    story.append(Spacer(1, 2))
    story.append(para(f"<b>BlackRock:</b> {clip(vendor['blackrock_connection'], 240)}", "body_sm"))

    others = [p for p in vendor["people"] if p["name"] != person["name"]]
    if others:
        story.append(Spacer(1, 2))
        who = "; ".join(f"{p['name']} ({p['title'].replace('&amp;', '&')})" for p in others[:4])
        story.append(para(f"<b>Others at {vendor['name']}:</b> {clip(who, 300)}", "body_sm"))

    # ---------- 05 private block (only in --with-notes builds) ----------
    if notes or engagements:
        story.append(Spacer(1, 7))
        story.append(section_head("Your Notes \u2014 Confidential", "05"))
        story.append(Spacer(1, 3))
        rows = []
        if notes:
            for label, key in [("Email", "email"), ("Phone", "phone"), ("Mobile", "mobile"),
                               ("Prefers", "preferred"), ("Assistant", "assistant"),
                               ("Location", "location"), ("In role since", "since"),
                               ("Relationship owner", "owner"), ("Last contact", "last_contact"),
                               ("Next touch", "next_touch")]:
                if notes.get(key):
                    rows.append(f"<b>{label}:</b> {notes[key]}")
            if notes.get("how_we_know"):
                rows.append(f"<b>How we know them:</b> {clip(notes['how_we_know'], 260)}")
            if notes.get("where_met"):
                rows.append(f"<b>Where we met:</b> {clip(notes['where_met'], 200)}")
            if notes.get("meeting_facts"):
                rows.append(f"<b>From meetings:</b> {clip(notes['meeting_facts'], 340)}")
            if notes.get("notes"):
                rows.append(f"<b>Notes:</b> {clip(notes['notes'], 320)}")
        if rows:
            story.append(para("<br/>".join(rows), "body_sm"))
        if engagements:
            story.append(Spacer(1, 3))
            story.append(para(f"<b>Live engagement with {vendor['name']}:</b>", "body_sm"))
            eng_lines = []
            for e in engagements[:4]:
                bits = [b for b in [e.get("portco"), e.get("use_case"), e.get("status")] if b]
                if bits:
                    eng_lines.append(" \u2014 ".join(bits))
            story.append(para(bullets([clip(x, 190) for x in eng_lines]) or "\u2014", "body_sm"))

    # ---------- footer ----------
    story.append(Spacer(1, 8))
    links = [f'<a href="{vendor["website"]}" color="#2B6684">{vendor["website"].replace("https://", "")}</a>']
    for l in (research or {}).get("links", [])[:3]:
        links.append(f'<a href="{l["url"]}" color="#2B6684">{l["label"]}</a>')
    if notes and notes.get("linkedin"):
        links.append(f'<a href="{notes["linkedin"]}" color="#2B6684">LinkedIn</a>')
    if vendor.get("linkedin"):
        links.append(f'<a href="{vendor["linkedin"]}" color="#2B6684">Company LinkedIn</a>')
    story.append(para(" &nbsp;\u00b7&nbsp; ".join(links), "src"))
    story.append(Spacer(1, 2))
    story.append(para(
        "Public professional and company information from the research layer of the GIP vendor tracker, "
        f"generated {BUILD_DATE}. Figures are as reported by the sources listed on the vendor's white page "
        "and are not independently audited. Research goes stale \u2014 re-run before external use."
        + (" This copy contains confidential GIP notes: do not circulate." if (notes or engagements) else ""),
        "src"))

    doc.build(story)


def load_workbook_data():
    """Read People and Engagements out of gip-tracker.xlsx. Returns ({}, {}) if absent."""
    book = os.path.join(ROOT_DIR, "gip-tracker.xlsx")
    if not os.path.exists(book):
        print("gip-tracker.xlsx not found — building public sheets only.")
        return {}, {}
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("openpyxl not installed — building public sheets only.")
        return {}, {}

    wb = load_workbook(book, data_only=True)
    people = {}
    if "People" in wb.sheetnames:
        ws = wb["People"]
        hdr = [c.value for c in ws[1]]

        def col(name):
            return hdr.index(name) if name in hdr else None

        idx = {k: col(v) for k, v in {
            "id": "Person ID", "linkedin": "LinkedIn URL", "email": "Work Email", "phone": "Phone",
            "location": "Location", "since": "In Role Since", "background": "Prior Background",
            "owner": "GIP Relationship Owner", "how_we_know": "How We Know Them",
            "last_contact": "Last Contact", "notes": "Notes / Facts",
            "mobile": "Mobile", "preferred": "Preferred Contact", "assistant": "Assistant / EA",
            "where_met": "Where We Met", "meeting_facts": "Facts From Meetings",
            "next_touch": "Next Touch"}.items()}
        for r in ws.iter_rows(min_row=2):
            if idx["id"] is None or not r[idx["id"]].value:
                continue
            rec = {}
            for k, i in idx.items():
                if i is not None and r[i].value not in (None, ""):
                    rec[k] = str(r[i].value).strip()
            pid = rec.pop("id", None)
            if pid and rec:
                people[pid] = rec

    engagements = {}
    if "Engagements" in wb.sheetnames:
        ws = wb["Engagements"]
        hdr = [c.value for c in ws[1]]

        def col(name):
            return hdr.index(name) if name in hdr else None

        iv, ip, iu, ist = col("Vendor"), col("PortCo Introduced"), col("Use Case"), col("Status")
        for r in ws.iter_rows(min_row=2):
            if iv is None or not r[iv].value:
                continue
            vendor = str(r[iv].value).strip()
            if vendor.upper().startswith("EXAMPLE"):
                continue
            row = {"portco": str(r[ip].value).strip() if ip is not None and r[ip].value else "",
                   "use_case": str(r[iu].value).strip() if iu is not None and r[iu].value else "",
                   "status": str(r[ist].value).strip() if ist is not None and r[ist].value else ""}
            if any(v.upper().startswith("EXAMPLE") for v in row.values()):
                continue
            engagements.setdefault(vendor.lower(), []).append(row)
    return people, engagements


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    with_notes = "--with-notes" in sys.argv

    people_notes, engagements = ({}, {})
    if with_notes:
        people_notes, engagements = load_workbook_data()
        out_dir = os.path.join(ROOT_DIR, "private-contact-sheets")
        os.makedirs(out_dir, exist_ok=True)
    else:
        out_dir = ROOT_DIR

    from pypdf import PdfReader

    made = 0
    for v in VENDORS:
        for p in v["people"]:
            pid = person_id(v["slug"], p["name"])
            if args and pid not in args and p["name"] not in args:
                continue
            out = os.path.join(out_dir, f"contact-{pid}-one-sheet.pdf")
            notes = people_notes.get(pid) if with_notes else None
            eng = engagements.get(v["name"].lower()) if with_notes else None
            for level in range(4):
                build_sheet(p, v, out, notes=notes, engagements=eng, level=level,
                            research=people_research.get(pid))
                if len(PdfReader(out).pages) == 1:
                    break
            made += 1

    where = "private-contact-sheets/" if with_notes else "the repo root"
    print(f"wrote {made} contact one-sheet(s) to {where}"
          + (" — confidential, gitignored" if with_notes else ""))


if __name__ == "__main__":
    main()
